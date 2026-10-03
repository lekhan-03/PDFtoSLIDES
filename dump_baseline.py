import sys
from pathlib import Path
from pptx import Presentation

def dump_pptx_text(pptx_path: Path, out_path: Path):
    prs = Presentation(str(pptx_path))
    lines = []
    for i, slide in enumerate(prs.slides):
        lines.append(f"--- Slide {i+1} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text_frame"):
                lines.append(shape.text.strip())
        lines.append("")
    out_path.write_text("\n".join(lines), encoding="utf-8")

def main():
    import subprocess
    import os
    
    baseline_dir = Path("tests/baseline")
    baseline_dir.mkdir(parents=True, exist_ok=True)
    
    files = [
        Path("data/tests/testres/Integrals.json"),
        Path("data/outputs/chemistry_2026_10_03_130427.json")
    ]
    
    for f in files:
        if not f.exists():
            print(f"File not found: {f}")
            continue
            
        print(f"Generating baseline for {f}...")
        pptx_path = baseline_dir / f.with_suffix(".pptx").name
        log_path = baseline_dir / f.with_suffix(".log").name
        txt_path = baseline_dir / f.with_suffix(".txt").name
        
        # Run generate_pptx.py
        cmd = [sys.executable, "-m", "src.generate_pptx", str(f), str(pptx_path)]
        res = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
        
        # Save logs
        log_path.write_text(res.stdout + "\n" + res.stderr, encoding="utf-8")
        
        # Dump text
        if pptx_path.exists():
            dump_pptx_text(pptx_path, txt_path)
            print(f"Saved {pptx_path.name} and {txt_path.name}")
        else:
            print(f"Failed to generate {pptx_path.name}")

if __name__ == "__main__":
    main()
