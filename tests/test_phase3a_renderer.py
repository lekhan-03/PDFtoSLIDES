"""Renderer ownership, script alignment, and source-traceability regressions."""
import html
import re
import zipfile

from docx import Document
from docx.oxml import parse_xml
from lxml import etree
from pptx import Presentation


def _math_xml(expression):
    return (
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        f"{expression}</m:oMath>"
    )


def test_phase3a_fallback_is_local_to_each_math_expression(monkeypatch):
    import src.generate_pptx as renderer
    import src.convert.math as math_converter

    monkeypatch.setattr(math_converter, "convert_math", lambda value: value)

    class FakePandoc:
        @staticmethod
        def convert_text(markdown, output_format, format, outputfile):
            expression = re.fullmatch(r"\$(.+)\$", markdown, flags=re.S).group(1)
            xml = (
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><w:body><w:p>'
                f'<m:oMath><m:r><m:t>{html.escape(expression)}</m:t></m:r></m:oMath>'
                '</w:p></w:body></w:document>'
            )
            with zipfile.ZipFile(outputfile, "w") as archive:
                archive.writestr("word/document.xml", xml)

    monkeypatch.setattr(renderer, "pypandoc", FakePandoc)
    presentation = Presentation()
    paragraph = presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_textbox(
        0, 0, 100, 100
    ).text_frame.paragraphs[0]
    diagnostics = []
    renderer._insert_omml_math_into_paragraph(
        paragraph, r"Electric field $E_{1}$ at distance $r$ from the charge", 14,
        renderer.OFF_WHITE, diagnostic_context={"question_id": "q-local", "component_id": "stem"},
        diagnostics=diagnostics,
    )

    xml = paragraph._p.xml
    root = etree.fromstring(xml.encode("utf-8"))
    assert "Electric field E₁ at distance r from the charge" in "".join(root.xpath(".//*[local-name()='t']/text()"))
    assert "AlternateContent" not in xml and "a14:m" not in xml
    assert [item["render_path"] for item in diagnostics] == ["compatibility", "compatibility"]


def test_math_uses_safe_compatibility_text_with_inherited_color(monkeypatch):
    import src.generate_pptx as renderer
    import src.convert.math as math_converter

    class FakePandoc:
        @staticmethod
        def convert_text(markdown, output_format, format, outputfile):
            expression = re.fullmatch(r"\$(.+)\$", markdown, flags=re.S).group(1)
            xml = (
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><w:body><w:p>'
                '<m:oMath><m:f><m:num><m:r><m:t>a</m:t></m:r></m:num>'
                '<m:den><m:r><m:t>b</m:t></m:r></m:den></m:f>'
                f'<m:r><m:t>{html.escape(expression)}</m:t></m:r></m:oMath>'
                '</w:p></w:body></w:document>'
            )
            with zipfile.ZipFile(outputfile, "w") as archive:
                archive.writestr("word/document.xml", xml)

    monkeypatch.setattr(renderer, "pypandoc", FakePandoc)
    monkeypatch.setattr(math_converter, "convert_math", lambda value: value)
    expressions = [
        "x² + y² = z²", r"\frac{a}{b}", "10²", "H₂O", "μ = 2πr",
        "Force = 10⁻⁴ N m", "x^2 and y^2", "a/b",
    ]
    prs = Presentation()
    paragraph = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
        0, 0, 100, 100
    ).text_frame.paragraphs[0]
    for expression in expressions:
        renderer._insert_omml_math_into_paragraph(
            paragraph, f"${expression}$", 16, renderer.OFF_WHITE
        )
    renderer._insert_omml_math_into_paragraph(
        paragraph, "$x^2$ + $y^2$", 16, renderer.GOLD
    )

    root = etree.fromstring(paragraph._p.xml.encode("utf-8"))
    namespaces = {
        "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
        "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
        "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    }
    assert not root.xpath(".//*[local-name()='AlternateContent' or local-name()='m']")
    assert "x² + y² = z²" in "".join(root.xpath(".//*[local-name()='t']/text()"))
    colors = root.xpath(".//*[local-name()='rPr']/*[local-name()='solidFill']/*[local-name()='srgbClr']/@val")
    assert colors[:len(expressions)] == ["F0F0F0"] * len(expressions)
    assert colors[len(expressions):] == ["FAD02C"] * 3


def test_native_math_styling_is_disabled_until_schema_verified():
    from src.generate_pptx import _style_native_math_color

    math = etree.fromstring(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<m:sSup><m:e><m:r><m:rPr><m:sty m:val="bi"/><m:ctrlPr><w:rPr>'
        '<w:b/><w:i/><w:color w:val="000000" w:themeColor="text1"/>'
        '</w:rPr></m:ctrlPr></m:rPr><m:t>x</m:t></m:r></m:e>'
        '<m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>'
        '<m:nary><m:naryPr><m:chr m:val="∫"/></m:naryPr><m:e><m:r><m:t>x</m:t>'
        '</m:r></m:e></m:nary></m:oMath>'
    )
    import pytest
    with pytest.raises(RuntimeError, match="deferred"):
        _style_native_math_color(math, (240, 240, 240))


def test_phase3a_unicode_script_alignment_matches_latex_compatibility():
    from src.generate_pptx import _inline_for_component

    nodes = [
        {"kind": "text", "value": "5.12 × 10", "source_block_id": 14, "source_order": 2},
        {"kind": "text", "value": "25", "vertical_align": "superscript",
         "source_block_id": 14, "source_order": 2},
    ]
    question = {
        "inline_content": nodes,
        "source_blocks": [{"block_id": "14.part_1", "role": "option"}],
    }
    aligned, source_id = _inline_for_component(question, r"5.12 × 10^{25}", "option:A")
    assert aligned is not None
    assert len(aligned) == 2
    assert source_id == 14


def test_phase3a_unmatched_option_diagnostic_uses_component_source():
    from src.generate_pptx import LayoutEngine

    question = {
        "id": "q-source", "question": "Stem", "options": {},
        "block_ids": [13, "14.part_1", "14.part_2", "15.part_1", "15.part_2"],
        "source_blocks": [
            {"block_id": 13, "role": "stem"},
            {"block_id": "14.part_1", "role": "option"},
            {"block_id": "14.part_2", "role": "option"},
            {"block_id": "15.part_1", "role": "option"},
            {"block_id": "15.part_2", "role": "option"},
        ],
        "inline_content": [
            {"kind": "text", "value": "original option A", "source_block_id": 14, "source_order": 2},
            {"kind": "text", "value": "original option B", "source_block_id": 14, "source_order": 2},
            {"kind": "text", "value": "original option C", "source_block_id": 15, "source_order": 3},
            {"kind": "text", "value": "original option D", "source_block_id": 15, "source_order": 3},
        ],
    }
    engine = LayoutEngine(Presentation(), question)
    expected = ["14.part_1", "14.part_2", "15.part_1", "15.part_2"]
    for letter in "ABCD":
        nodes, source_id = engine._component_inline(f"unmatched compatibility option {letter}", f"option:{letter}")
        assert nodes is None
        assert source_id == expected[ord(letter) - ord("A")]
    assert [item["source_block_id"] for item in engine.render_diagnostics] == expected
    assert all(item["node_id"] for item in engine.render_diagnostics)


def test_phase3a_canonical_script_node_is_emitted_once():
    import src.generate_pptx as renderer

    presentation = Presentation()
    paragraph = presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_textbox(
        0, 0, 100, 100
    ).text_frame.paragraphs[0]
    renderer._render_canonical_inline(
        paragraph,
        [
            {"kind": "text", "value": "5.12 × 10", "source_block_id": "b-opt"},
            {"kind": "text", "value": "25", "vertical_align": "superscript", "source_block_id": "b-opt"},
        ],
        14, renderer.OFF_WHITE, False, "Cambria", [],
        {"question_id": "q-once", "component_id": "option:A"},
    )
    texts = paragraph._p.xpath(".//*[local-name()='t']/text()")
    assert "5.12 × 10" in texts
    assert texts.count("25") == 1
    assert paragraph._p.xml.count('baseline="30000"') == 1


def test_safe_pptx_package_contains_only_well_formed_text_math(tmp_path):
    import json
    import posixpath
    from src.generate_pptx import generate

    payload = {"questions": [{
        "id": "safe-q1", "question": "Value: $x_B = n_A + n_B$", "options": {"A": "x+1"},
        "inline_content": [
            {"kind": "text", "value": "Value: ", "source_block_id": "stem-1"},
            {"kind": "math_sequence", "value": "$x_B = n_A + n_B$", "source_block_id": "stem-1"},
        ],
        "source_blocks": [{"block_id": "stem-1", "role": "stem"}],
    }]}
    source, target = tmp_path / "safe.json", tmp_path / "safe.pptx"
    source.write_text(json.dumps(payload), encoding="utf-8")
    result = generate(source, target)
    assert result["rendered_questions"] == 1
    with zipfile.ZipFile(target) as archive:
        assert archive.testzip() is None
        names = set(archive.namelist())
        for name in names:
            if name.endswith(".xml") or name.endswith(".rels"):
                root = etree.fromstring(archive.read(name))
                assert not root.xpath(".//*[local-name()='AlternateContent']")
                if name.endswith(".rels"):
                    normalized_name = name.replace("\\", "/")
                    source_part = (normalized_name.replace("/_rels/", "/").removesuffix(".rels")
                                   if "/_rels/" in normalized_name else "")
                    base = posixpath.dirname(source_part)
                    for rel in root.xpath("//*[local-name()='Relationship']"):
                        if rel.get("TargetMode") == "External":
                            continue
                        target_name = posixpath.normpath(posixpath.join(base, rel.get("Target")))
                        assert target_name in names, (name, target_name)
        slide = etree.fromstring(archive.read("ppt/slides/slide2.xml"))
        text = "".join(slide.xpath(".//*[local-name()='t']/text()"))
        assert "Value: x_B = n_A + n_B" in text
        assert not slide.xpath(".//*[namespace-uri()='http://schemas.microsoft.com/office/drawing/2010/main']")
    Presentation(str(target))


def test_canonical_english_text_is_never_reclassified_as_math():
    import src.generate_pptx as renderer

    phrases = ["explanation", "expression", "expansion", "experiment"]
    prs = Presentation()
    paragraph = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
        0, 0, 100, 100
    ).text_frame.paragraphs[0]
    renderer._render_canonical_inline(
        paragraph, [{"kind": "text", "value": " | ".join(phrases), "source_block_id": "b-text"}],
        14, renderer.GOLD, False, "Cambria", [], {"component_id": "question"},
    )
    assert "".join(paragraph._p.xpath(".//*[local-name()='t']/text()")) == " | ".join(phrases)
    assert "AlternateContent" not in paragraph._p.xml


def test_phase3a_malformed_math_fallback_is_readable_and_traceable(monkeypatch):
    import src.generate_pptx as renderer

    monkeypatch.setattr(renderer, "pypandoc", None)
    presentation = Presentation()
    paragraph = presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_textbox(
        0, 0, 100, 100
    ).text_frame.paragraphs[0]
    diagnostics = []
    malformed = r"$\sqrt(3)\times10^{-4}$ tail \_{eq}\_{}"
    renderer._insert_omml_math_into_paragraph(
        paragraph, malformed, 14, renderer.OFF_WHITE,
        diagnostic_context={"question_id": "q-bad-math", "component_id": "stem", "source_block_id": "b-1"},
        diagnostics=diagnostics,
    )
    text = "".join(paragraph._p.xpath(".//*[local-name()='t']/text()"))
    assert "sqrt(3)" not in text
    assert "_{}" not in text and "_{eq}" not in text and "$" not in text
    assert "√(3)" in text
    assert any(item["code"] == "MATH_FALLBACK_USED" for item in diagnostics)
    assert all(item["question_id"] == "q-bad-math" and item.get("original") for item in diagnostics)


def test_phase3a_fraction_fallback_preserves_fraction_meaning():
    from src.generate_pptx import _fallback_math_text

    assert _fallback_math_text(r"\frac{mω^{2}x}{e}") == "(mω² x)/(e)"


def test_phase3a_isolated_operator_text_does_not_create_math_node(monkeypatch):
    import src.generate_pptx as renderer

    monkeypatch.setattr(renderer, "render_math_text", lambda value: f"${value}$")
    presentation = Presentation()
    paragraph = presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_textbox(
        0, 0, 100, 100
    ).text_frame.paragraphs[0]
    diagnostics = []
    renderer._render_canonical_inline(
        paragraph, [{"kind": "text", "value": "×", "source_block_id": "b-1"}], 14,
        renderer.OFF_WHITE, False, "Cambria", diagnostics, {"component_id": "stem"},
    )
    assert "<a14:m" not in paragraph._p.xml
    assert paragraph._p.xpath(".//*[local-name()='t']/text()") == ["×"]
    assert diagnostics == []


def test_phase3a_explicit_math_does_not_reparse_adjacent_component_text(monkeypatch):
    import src.generate_pptx as renderer
    import src.convert.math as math_converter

    monkeypatch.setattr(math_converter, "convert_math", lambda value: value)
    monkeypatch.setattr(renderer, "pypandoc", None)
    monkeypatch.setattr(renderer, "render_math_text", lambda value: (_ for _ in ()).throw(
        AssertionError("explicit math components must not be auto-reparsed")
    ))
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    diagnostics = []
    shape = renderer._tb(
        slide, r"A copper rod AB rotates along x; option a is $\frac{mω^{2}x}{e}$",
        0, 0, 600, 100, 14, renderer.OFF_WHITE,
        diagnostics=diagnostics, diagnostic_context={"question_id": "q-duplicate", "component_id": "stem"},
    )
    emitted_text = "".join(shape.text_frame.paragraphs[0]._p.xpath(".//*[local-name()='t']/text()"))
    assert emitted_text.count("AB") == 1
    assert "mω² x" in emitted_text and "(e)" in emitted_text
    assert not any("component text" in str(item.get("reason")) for item in diagnostics)


def test_phase3a_math_count_summary_reconciles_native_and_fallback(tmp_path):
    import json
    from src.generate_pptx import generate

    payload = {"questions": [{
        "id": "q-count", "question": "$x^{2}$", "options": {},
        "inline_content": [{
            "kind": "math_sequence", "value": "$x^{2}$", "source_block_id": "b-1",
            "source": {"omml": _math_xml("<m:r><m:t>x</m:t></m:r>")},
        }],
    }]}
    source, target = tmp_path / "math_count.json", tmp_path / "math_count.pptx"
    source.write_text(json.dumps(payload), encoding="utf-8")
    result = generate(source, target)
    counts = result["math_render_counts"]
    assert counts["canonical_math_nodes"] == 1
    assert counts["native_math_entries"] == 0
    assert counts["fallback_math_entries"] == 0
    assert counts["compatibility_math_nodes"] == 1
    assert counts["failed_math_nodes"] == 0
    outcomes = result["math_render_outcomes"]
    assert len(outcomes) == 1
    assert outcomes[0]["math_id"].startswith("Q-COUNT-BLOCK-B-1-MATH-")
    assert outcomes[0]["render_path"] == "compatibility"
    assert counts["unrendered_canonical_math_nodes"] == 0
    with zipfile.ZipFile(target) as archive:
        xml = archive.read("ppt/slides/slide2.xml").decode("utf-8")
    assert "AlternateContent" not in xml and "<a14:m" not in xml


def test_phase3a_existing_question_component_mismatch_stays_traceable(tmp_path):
    import json
    from pathlib import Path
    from lxml import etree
    from src.generate_pptx import generate

    fixture = Path(__file__).resolve().parents[1] / "data" / "tests" / "testing_file.json"
    question = next(item for item in json.loads(fixture.read_text(encoding="utf-8"))["questions"]
                    if item.get("id") == 16)
    assert question["options"] == {}
    assert "a)" in question["question"] and any(
        block.get("role") == "option" for block in question.get("source_blocks", [])
    )
    source = tmp_path / "known_mismatch.json"
    target = tmp_path / "known_mismatch.pptx"
    original = {key: value for key, value in json.loads(json.dumps(question)).items() if key != "render"}
    source.write_text(json.dumps({"questions": [question]}), encoding="utf-8")
    result = generate(source, target)
    saved = json.loads(source.read_text(encoding="utf-8"))["questions"][0]
    assert {k: v for k, v in saved.items() if k != "render"} == original
    diagnostic = next(item for item in result["render_diagnostics"]
                      if item.get("component_id") == "question")
    assert diagnostic["code"] == "TEXT_LAYOUT_DEGRADED"
    assert diagnostic["source_block_id"] == "61"
    assert diagnostic["node_id"] == "61:61:1"
    with zipfile.ZipFile(target) as archive:
        xml = etree.fromstring(archive.read("ppt/slides/slide2.xml"))
    assert not xml.xpath(".//*[local-name()='AlternateContent']")
    assert not xml.xpath(".//*[local-name()='a14:m']")
    slide_text = "".join(xml.xpath(".//*[local-name()='t']/text()"))
    assert "copper rod" in slide_text
