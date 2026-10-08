"""Phase 3 math, canonical-node, table, diagnostics, and PDF regressions."""
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor


def _math_xml(expression):
    return (
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        f'{expression}</m:oMath>'
    )


def _math_run(text):
    return f'<m:r><m:t>{text}</m:t></m:r>'


def test_ph3_math_001_golden_unicode_and_latex_semantics():
    from src.formula_render import auto_latex

    expected = {
        "10⁵": r"$10^{5}$", "10⁻⁵": r"$10^{-5}$", "x²": r"$x^{2}$",
        "x₁": r"$x_{1}$", "x₁²": r"$x_{1}^{2}$", "C⁻¹": r"$C^{-1}$",
        "mL⁻¹": r"$mL^{-1}$", "ε₀": r"$\varepsilon_{0}$",
        "μ": r"$\mu$", "π": r"$\pi$", "√3": r"$\sqrt{3}$",
        "√(a²+b²)": r"$\sqrt{a^{2}+b^{2}}$", "Na⁺": r"$Na^{+}$",
        "Cl⁻": r"$Cl^{-}$", "SO₄²⁻": r"$SO_{4}^{2-}$",
        "H₂O": r"$H_{2}O$", "CO₂": r"$CO_{2}$",
        r"\vec{F}=q\vec{E}": r"$\vec{F}=q\vec{E}$",
    }
    for source, semantic in expected.items():
        assert auto_latex(source) == semantic, source
    for symbol in ("×", "±", "90°", "45°", "180°", "a/b"):
        assert auto_latex(symbol) == symbol


def test_math_prompt_regressions_preserve_q1_q3_and_plain_english():
    from src.formula_render import auto_latex, latex_to_readable

    q1 = "∆mixH = 0 and ∆mixV = 0"
    q3 = "x_B = n_A / (n_A + n_B)"
    assert auto_latex(q1) == q1
    assert auto_latex(q3).count("$") % 2 == 0
    assert latex_to_readable(q3) == q3
    for plain_text in ("explanation", "experimental", "experience"):
        assert auto_latex(plain_text) == plain_text
        assert not re.search(r"\\exp", auto_latex(plain_text))
    assert "+" in auto_latex("n_A + n_B")


def test_docx_math_operator_vertical_align_does_not_turn_plus_into_subscript(tmp_path):
    from src.extract_docx import extract_docx_blocks

    doc = Document()
    paragraph = doc.add_paragraph("Formula: ")
    paragraph._p.append(parse_xml(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<m:f><m:num><m:r><m:t>n</m:t></m:r><m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e>'
        '<m:sub><m:r><m:t>A</m:t></m:r></m:sub></m:sSub></m:num><m:den>'
        '<m:r><m:t>n</m:t></m:r><m:sSub><m:e><m:r><m:t>A</m:t></m:r></m:e>'
        '<m:sub><m:r><m:t>A</m:t></m:r></m:sub></m:sSub>'
        '<m:r><w:rPr><w:vertAlign w:val="subscript"/></w:rPr><m:t>+</m:t></m:r>'
        '<m:r><m:t>n</m:t></m:r><m:sSub><m:e><m:r><m:t>B</m:t></m:r></m:e>'
        '<m:sub><m:r><m:t>B</m:t></m:r></m:sub></m:sSub>'
        '</m:den></m:f></m:oMath>'
    ))
    path = tmp_path / "operator_vertical_align.docx"
    doc.save(path)
    blocks = extract_docx_blocks(str(path))
    math_node = next(node for node in blocks[0].inline_content if node.kind == "math_sequence")
    assert "₊" not in math_node.value
    assert "+" in math_node.value
    source_root = etree.fromstring(math_node.source["omml"].encode("utf-8"))
    plus_run = source_root.xpath(".//*[local-name()='r'][.//*[local-name()='t' and text()='+']]")
    assert len(plus_run) == 1
    assert not plus_run[0].xpath(".//*[local-name()='vertAlign']")


def test_canonical_text_nodes_are_never_reclassified_as_math(monkeypatch):
    import src.generate_pptx as renderer

    monkeypatch.setattr(renderer, "render_math_text", lambda _value: (_ for _ in ()).throw(
        AssertionError("canonical text must not be reclassified")))
    prs = Presentation()
    paragraph = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
        0, 0, 100, 100
    ).text_frame.paragraphs[0]
    diagnostics = []
    renderer._render_canonical_inline(
        paragraph, [{"kind": "text", "value": "explanation"}], 16,
        renderer.OFF_WHITE, False, "Cambria", diagnostics, {"component_id": "stem"},
    )
    assert "<a14:m" not in paragraph._p.xml
    assert paragraph._p.xpath(".//*[local-name()='t']/text()") == ["explanation"]
    assert diagnostics == []


def test_math_inherits_stem_option_and_table_base_colors():
    from lxml import etree
    import src.generate_pptx as renderer
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    nodes = [{"kind": "math_sequence", "value": "$x^2$", "source_block_id": "b-math"}]
    expected = [("stem", renderer.GOLD), ("option", renderer.CYAN), ("table", renderer.OFF_WHITE)]
    for index, (component, color) in enumerate(expected):
        shape = slide.shapes.add_textbox(0, index * 100, 100, 100)
        paragraph = shape.text_frame.paragraphs[0]
        renderer._render_canonical_inline(paragraph, nodes, 16, color, False, "Cambria", [],
                                          {"component_id": component})
        xml = etree.fromstring(paragraph._p.xml.encode("utf-8"))
        assert xml.xpath(".//*[local-name()='solidFill']/*[local-name()='srgbClr']/@val")
        assert "AlternateContent" not in paragraph._p.xml
        color_hex = "".join(f"{byte:02X}" for byte in color)
        assert set(xml.xpath(".//*[local-name()='solidFill']/*[local-name()='srgbClr']/@val")) == {color_hex}

    option = slide.shapes.add_textbox(0, 400, 100, 100)
    renderer._shape_text(option, slide, "A", "$x^2$", 16, 16)
    body_run = option.text_frame.paragraphs[0].runs[1]
    assert body_run.font.color.rgb == renderer.CYAN


def test_ph3_math_002_latex_to_readable_pdf_text_preserves_scripts_symbols():
    from src.formula_render import latex_to_readable

    cases = {
        "10⁵ × N C⁻¹": "10⁵ × N C⁻¹",
        "ε₀ E": "ε₀ E",
        "SO₄²⁻": "SO₄²⁻",
        "√(a²+b²)": "√(a²+b²)",
        "90° ± 5°": "90° ± 5°",
        r"\vec{F}=q\vec{E}": "F⃗=qE⃗",
        r"\frac{a+b}{c+d}": "(a+b)/(c+d)",
    }
    for source, expected in cases.items():
        assert latex_to_readable(source) == expected, source


def test_ph3_render_001_source_omml_fraction_uses_safe_semantic_text():
    from src.generate_pptx import _render_canonical_inline

    raw = _math_xml(
        '<m:f><m:num><m:r><m:t>a+b</m:t></m:r></m:num>'
        '<m:den><m:r><m:t>c+d</m:t></m:r></m:den></m:f>'
    )
    node = {"kind": "math_sequence", "value": r"$\frac{a+b}{c+d}$",
            "source": {"omml_tag": "oMath", "omml": raw}, "children": [
                {"kind": "fraction", "value": r"\frac{a+b}{c+d}", "children": [
                    {"kind": "numerator", "value": "a+b", "children": [{"kind": "math_sequence", "value": "a+b"}]},
                    {"kind": "denominator", "value": "c+d", "children": [{"kind": "math_sequence", "value": "c+d"}]},
                ]}], "source_block_id": "b-frac"}
    prs = Presentation()
    paragraph = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, 100, 100).text_frame.paragraphs[0]
    diagnostics = []
    _render_canonical_inline(paragraph, [node], 16, (255, 255, 255), False, "Cambria", diagnostics,
                             {"question_id": "q-frac", "component_id": "question"})

    xml = paragraph._p.xml
    assert "AlternateContent" not in xml and "<a14:m" not in xml and "<m:f" not in xml
    assert "(a+b)/(c+d)" in xml
    assert diagnostics[0]["code"] == "MATH_COMPATIBILITY_RENDERED"


def test_ph3_render_002_formatted_runs_and_breaks_are_native_nodes():
    from src.generate_pptx import _render_canonical_inline

    nodes = [
        {"kind": "text", "value": "H", "formatting": {"bold": True}},
        {"kind": "text", "value": "2", "vertical_align": "subscript"},
        {"kind": "text", "value": "O"},
        {"kind": "line_break", "value": ""},
        {"kind": "text", "value": "x", "formatting": {"italic": True}},
        {"kind": "text", "value": "2", "vertical_align": "superscript"},
    ]
    prs = Presentation()
    paragraph = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, 100, 100).text_frame.paragraphs[0]
    diagnostics = []
    _render_canonical_inline(paragraph, nodes, 16, RGBColor(255, 255, 255), False, "Cambria", diagnostics,
                             {"question_id": "q-rich", "component_id": "question"})
    xml = paragraph._p.xml
    assert "baseline=\"-25000\"" in xml
    assert "baseline=\"30000\"" in xml
    assert "<a:br" in xml
    assert "b=\"1\"" in xml and "i=\"1\"" in xml
    assert diagnostics == []


def test_ph3_diag_001_unknown_node_and_math_fallback_are_question_scoped(monkeypatch):
    import src.generate_pptx as renderer

    monkeypatch.setattr(renderer, "pypandoc", None)
    prs = Presentation()
    paragraph = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, 100, 100).text_frame.paragraphs[0]
    diagnostics = []
    context = {"question_id": "q-diag", "component_id": "option:B", "source_block_id": "b-2"}
    renderer._insert_omml_math_into_paragraph(paragraph, r"$x^2$", 14, renderer.OFF_WHITE,
                                              diagnostic_context={**context, "node_id": "math-1"},
                                              diagnostics=diagnostics)
    renderer._render_canonical_inline(paragraph, [{"kind": "unknown_widget", "value": "widget",
                                                    "source_block_id": "b-3"}], 14,
                                       renderer.OFF_WHITE, False, "Cambria", diagnostics, context)
    assert {item["code"] for item in diagnostics} == {"MATH_COMPATIBILITY_RENDERED", "RENDER_NODE_UNSUPPORTED"}
    assert all(item["question_id"] == "q-diag" for item in diagnostics)
    assert all(item["component_id"] == "option:B" for item in diagnostics)
    assert all(item["source_block_id"] in ("b-2", "b-3") for item in diagnostics)


def test_ph3_render_003_json_component_renders_canonical_omml(tmp_path):
    from src.generate_pptx import generate

    raw = _math_xml('<m:sSup><m:e>' + _math_run("x") + '</m:e><m:sup>' + _math_run("2") + '</m:sup></m:sSup>')
    payload = {"questions": [{
        "id": "q-inline", "question": r"$x^{2}$", "options": {},
        "inline_content": [{"kind": "math_sequence", "value": r"$x^{2}$",
                            "source_block_id": "b-1", "source_order": 1,
                            "source": {"omml_tag": "oMath", "omml": raw}, "children": []}],
        "source_blocks": [{"block_id": "b-1", "role": "stem"}],
    }]}
    source = tmp_path / "canonical.json"
    out = tmp_path / "canonical.pptx"
    source.write_text(json.dumps(payload), encoding="utf-8")
    result = generate(source, out)
    with zipfile.ZipFile(out) as archive:
        slide_xml = archive.read("ppt/slides/slide2.xml").decode("utf-8")
    assert "AlternateContent" not in slide_xml and "<a14:m" not in slide_xml
    assert result["math_render_counts"]["compatibility_math_nodes"] == 1
    saved = json.loads(source.read_text(encoding="utf-8"))
    assert saved["questions"][0]["render"]["status"] == "rendered"


def test_ph3_table_001_docx_table_rich_math_and_scripts_survive_json_to_pptx(tmp_path):
    from src.extract_docx import extract_docx_blocks
    from src.parser_v2 import reconstruct_components
    from src.generate_pptx import generate

    doc = Document()
    p = doc.add_paragraph("1. Read the table and select the correct row.")
    p.runs[0].bold = True
    table = doc.add_table(rows=1, cols=3)
    table.cell(0, 0).text = "Quantity"
    math_p = table.cell(0, 1).paragraphs[0]
    math_p._p.append(parse_xml(_math_xml(
        '<m:f><m:num>' + _math_run("a+b") + '</m:num><m:den>' + _math_run("c+d") + '</m:den></m:f>'
    )))
    chem_p = table.cell(0, 2).paragraphs[0]
    chem_p.add_run("SO")
    sub = chem_p.add_run("4")
    sub.font.subscript = True
    sup = chem_p.add_run("2−")
    sup.font.superscript = True
    docx_path = tmp_path / "rich_table.docx"
    doc.save(docx_path)

    blocks = extract_docx_blocks(str(docx_path))
    question = reconstruct_components({"blocks": blocks, "section": "Math"})
    question["id"] = "q-table-rich"
    assert question["table"]["rows"][0] == ["Quantity", r"$\frac{a+b}{c+d}$", "SO₄²⁻"]
    assert any(item.vertical_align == "subscript" for item in blocks[1].children[0].children[2].children[0].inline_content)
    assert any(item.vertical_align == "superscript" for item in blocks[1].children[0].children[2].children[0].inline_content)
    assert question["table"]["cell_content"][0][1]["inline_content"][0]["kind"] == "math_sequence"
    json_path, pptx_path = tmp_path / "table.json", tmp_path / "table.pptx"
    json_path.write_text(json.dumps({"questions": [question]}), encoding="utf-8")
    result = generate(json_path, pptx_path)
    with zipfile.ZipFile(pptx_path) as archive:
        slide_xml = archive.read("ppt/slides/slide2.xml").decode("utf-8")
    assert "AlternateContent" not in slide_xml and "<a14:m" not in slide_xml
    assert "(a+b)/(c+d)" in slide_xml
    assert 'baseline="-25000"' in slide_xml and 'baseline="30000"' in slide_xml
    assert all(item["severity"] == "info" for item in result["render_diagnostics"])


def test_ph3_pdf_001_json_to_pdf_keeps_math_chemistry_degree_and_table(tmp_path):
    import pymupdf
    from src.generate_pdf import generate_pdf

    payload = {"questions": [{
        "id": "q-pdf", "source_question_number": "1", "section": "Physics",
        "question": "10⁵ × N C⁻¹ and ε₀ E; SO₄²⁻ at 90° ± 5°; √(a²+b²).",
        "options": {"A": "Na⁺", "B": "Cl⁻", "C": "μ + π", "D": "45°"},
        "table": {"rows": [["Quantity", r"\frac{a+b}{c+d}", "H₂O"]]},
    }]}
    json_path, pdf_path = tmp_path / "pdf.json", tmp_path / "math.pdf"
    json_path.write_text(json.dumps(payload), encoding="utf-8")
    result = generate_pdf(str(json_path), str(pdf_path))
    pdf = pymupdf.open(pdf_path)
    text = "\n".join(page.get_text() for page in pdf)
    pdf.close()
    for expected in ("10⁵", "×", "C⁻¹", "ε₀", "SO₄²⁻", "90°", "±", "Na⁺", "Cl⁻", "μ", "π", "45°", "H₂O"):
        assert expected in text
    assert "\\frac" not in text and "$" not in text
    assert result["rendered_questions"] == 1

