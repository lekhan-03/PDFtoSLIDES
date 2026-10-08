import json
from pptx import Presentation
p=json.load(open('SolutionNEET_safe.json',encoding='utf-8'))
prs=Presentation('SolutionNEET_safe.pptx')
missing=[]; dup=[]
for i,q in enumerate(p['questions'],start=1):
    slide=prs.slides[i]
    texts=[s.text for s in slide.shapes if getattr(s,'has_text_frame',False)]
    for letter in q.get('options',{}):
        matches=[t for t in texts if t.lstrip().startswith(letter+'.')]
        if not matches: missing.append((i,letter))
        if len(matches)>1: dup.append((i,letter,len(matches)))
print('option_labels_missing',missing)
print('option_labels_duplicated',dup)
print('per_question_option_nodes',sum(len(q.get('options',{})) for q in p['questions']))
