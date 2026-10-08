from pptx import Presentation
p=Presentation('SolutionNEET_safe.pptx')
for i in [2,3]:
 print('slide',i)
 for s in p.slides[i].shapes:
  if getattr(s,'has_text_frame',False) and s.text.strip(): print(repr(s.text[:200]))
