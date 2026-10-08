import json,re
from pptx import Presentation
p=json.load(open('SolutionNEET_safe.json',encoding='utf-8'))
prs=Presentation('SolutionNEET_safe.pptx')
norm=lambda s: re.sub(r'\s+','',s).casefold()
missing=[]; dup=[]
for i,q in enumerate(p['questions'],start=1):
    texts=[s.text for s in prs.slides[i].shapes if getattr(s,'has_text_frame',False)]
    for letter,body in q.get('options',{}).items():
        expected=norm(body)
        matches=[t for t in texts if t.startswith(letter+'.') and norm(t.split('.',1)[1])==expected]
        if not matches: missing.append((i,letter,body))
        if len(matches)>1: dup.append((i,letter,len(matches)))
print('option_contents_missing',len(missing),missing[:8])
print('option_contents_duplicated',len(dup),dup[:8])
