# ── LayoutEngine Architecture ──

import argparse
import json
import re
import sys
import os
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Inches, Pt, Emu
except ImportError:
    pass

try:
    from src.formula_render import render_math_text
except ImportError:
    render_math_text = None


BG         = RGBColor(0x05, 0x05, 0x05)
GOLD       = RGBColor(0xFA, 0xD0, 0x2C)
CYAN       = RGBColor(0x00, 0xB0, 0xF0)
OFF_WHITE  = RGBColor(0xF0, 0xF0, 0xF0)
DIM_WHITE  = RGBColor(0xAA, 0xAA, 0xAA)
BADGE_BG   = RGBColor(0x1A, 0x16, 0x08)
LOGO_BG    = RGBColor(0x1C, 0x1C, 0x38)
DIVIDER_LINE = RGBColor(0x33, 0x33, 0x66)

SW = 13.33
SH = 7.5
M = 0.5
Q_L = M
Q_TOP = 0.80
Q_W = SW - 2 * M
MAX_Y = SH - 0.4  # Footer margin

class LayoutEngine:
    """Handles flowing content vertically, and spawning continuation slides if needed."""
    def __init__(self, prs, q, logo_path=None):
        self.prs = prs
        self.q = q
        self.logo_path = logo_path
        self.slide = None
        self.y = Q_TOP
        self.slide_count = 0
        self.new_slide()
        
    def new_slide(self):
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        
        # Background
        fill = self.slide.background.fill
        fill.solid()
        fill.fore_color.rgb = BG
        
        # Draw Logo
        LOGO_W = 1.5
        LOGO_H = 0.5
        LOGO_L = SW - M - LOGO_W
        HDR_TOP = 0.20
        if self.logo_path and self.logo_path.exists():
            self.slide.shapes.add_picture(str(self.logo_path), Inches(LOGO_L), Inches(HDR_TOP - 0.1), width=Inches(LOGO_W))
        
        # Citation Badge (only on first slide)
        if self.slide_count == 0:
            apps = self.q.get("appearances", [])
            if apps:
                # Merge multiple appearances
                citation = ", ".join([a.strip("()[]") for a in apps])
                self._draw_badge(citation)
        else:
            # Continuation marker
            self._draw_badge("Continued")
            
        self.y = Q_TOP
        self.slide_count += 1
        
    def _draw_badge(self, text):
        from src.generate_pptx import _calc_badge_width, _rect
        w = _calc_badge_width(text, 15.0)
        _rect(self.slide, M, 0.20, w, 0.44, fill=BADGE_BG, rounded=True, border=GOLD)
        _rect(self.slide, M + 0.03, 0.20 + 0.05, 0.06, 0.44 - 0.10, fill=GOLD, rounded=False)
        txb = self.slide.shapes.add_textbox(Inches(M + 0.15), Inches(0.20), Inches(w - 0.18), Inches(0.44))
        tf = txb.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = text
        r.font.size = Pt(15)
        r.font.color.rgb = GOLD
        r.font.bold = True
        r.font.name = 'Cambria'
        
    def space_available(self, required_height):
        return (self.y + required_height) <= MAX_Y
        
    def ensure_space(self, required_height):
        if not self.space_available(required_height) and self.y > Q_TOP:
            self.new_slide()
            
    def add_text(self, text, font_size=18, color=GOLD, bold=True):
        from src.generate_pptx import _get_content_height, _tb
        if not text:
            return
        h = _get_content_height(text, Q_W, font_size)
        self.ensure_space(h)
        _tb(self.slide, text, Q_L, self.y, Q_W, h, size=font_size, color=color, bold=bold)
        self.y += h + 0.1
        
    def add_image(self, img_path):
        from pptx.util import Inches
        if not os.path.exists(img_path):
            print(f"Warning: Image {img_path} not found")
            return
            
        from PIL import Image
        try:
            with Image.open(img_path) as img:
                w_px, h_px = img.size
        except Exception:
            print(f"Warning: Failed to read image {img_path}")
            return
            
        aspect = h_px / w_px
        
        # Max dimensions for image
        MAX_IMG_W = Q_W * 0.8
        MAX_IMG_H = 3.5
        
        # Fit to max constraints
        target_w = MAX_IMG_W
        target_h = target_w * aspect
        
        if target_h > MAX_IMG_H:
            target_h = MAX_IMG_H
            target_w = target_h / aspect
            
        self.ensure_space(target_h)
        self.slide.shapes.add_picture(img_path, Inches(Q_L), Inches(self.y), width=Inches(target_w), height=Inches(target_h))
        self.y += target_h + 0.1
        
    def add_options(self, opts):
        if not opts:
            return
            
        from src.generate_pptx import _can_fit_1x4, _get_content_height, OPT_W_1X4, GAP_1X4, OPT_W_2X2, GAP_2X2, GAP_V_2X2, OPT_TEXT_W_1X4, OPT_TEXT_W_2X2, _shape_text
        
        # Try 1x4
        if _can_fit_1x4(opts):
            h_a = _get_content_height(opts.get('A', ''), OPT_TEXT_W_1X4, 16)
            h_b = _get_content_height(opts.get('B', ''), OPT_TEXT_W_1X4, 16)
            h_c = _get_content_height(opts.get('C', ''), OPT_TEXT_W_1X4, 16)
            h_d = _get_content_height(opts.get('D', ''), OPT_TEXT_W_1X4, 16)
            row_h = max(0.50, h_a, h_b, h_c, h_d)
            
            self.ensure_space(row_h)
            for i, letter in enumerate("ABCD"):
                l = M + i * (OPT_W_1X4 + GAP_1X4)
                txb = self.slide.shapes.add_textbox(Inches(l), Inches(self.y), Inches(OPT_W_1X4), Inches(row_h))
                _shape_text(txb, self.slide, letter, opts.get(letter, ""),
                            letter_size=16, body_size=16,
                            left_in=l, top_in=self.y, width_in=OPT_W_1X4)
            self.y += row_h + 0.1
        else:
            # Try 2x2
            h_a = _get_content_height(opts.get('A', ''), OPT_TEXT_W_2X2, 16)
            h_b = _get_content_height(opts.get('B', ''), OPT_TEXT_W_2X2, 16)
            row1_h = max(0.50, h_a, h_b)
            
            h_c = _get_content_height(opts.get('C', ''), OPT_TEXT_W_2X2, 16)
            h_d = _get_content_height(opts.get('D', ''), OPT_TEXT_W_2X2, 16)
            row2_h = max(0.50, h_c, h_d)
            
            # If 2x2 doesn't fit on one slide, we might need to split it!
            # Let's ensure space for row 1
            self.ensure_space(row1_h)
            col_left = M
            col_right = M + OPT_W_2X2 + GAP_2X2
            
            # Row 1
            txb_a = self.slide.shapes.add_textbox(Inches(col_left), Inches(self.y), Inches(OPT_W_2X2), Inches(row1_h))
            _shape_text(txb_a, self.slide, "A", opts.get("A", ""), 16, 16, left_in=col_left, top_in=self.y, width_in=OPT_W_2X2)
            if "B" in opts:
                txb_b = self.slide.shapes.add_textbox(Inches(col_right), Inches(self.y), Inches(OPT_W_2X2), Inches(row1_h))
                _shape_text(txb_b, self.slide, "B", opts.get("B", ""), 16, 16, left_in=col_right, top_in=self.y, width_in=OPT_W_2X2)
            
            self.y += row1_h + GAP_V_2X2
            
            # Row 2
            if "C" in opts or "D" in opts:
                self.ensure_space(row2_h)
                if "C" in opts:
                    txb_c = self.slide.shapes.add_textbox(Inches(col_left), Inches(self.y), Inches(OPT_W_2X2), Inches(row2_h))
                    _shape_text(txb_c, self.slide, "C", opts.get("C", ""), 16, 16, left_in=col_left, top_in=self.y, width_in=OPT_W_2X2)
                if "D" in opts:
                    txb_d = self.slide.shapes.add_textbox(Inches(col_right), Inches(self.y), Inches(OPT_W_2X2), Inches(row2_h))
                    _shape_text(txb_d, self.slide, "D", opts.get("D", ""), 16, 16, left_in=col_right, top_in=self.y, width_in=OPT_W_2X2)
                
                self.y += row2_h + 0.1

    def add_table(self, table_data):
        # ... logic to render table using PPTX tables ...
        pass
